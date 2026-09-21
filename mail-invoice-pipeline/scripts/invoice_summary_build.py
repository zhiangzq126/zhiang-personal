# -*- coding: utf-8 -*-
"""发票汇总：读取 <数据目录>/invoices 下的发票 PDF -> 提取字段 -> 校验去重 -> 生成报销汇总表 xlsx

数据目录 = 环境变量 INVOICE_WORKDIR，未设置时为当前目录；
汇总表输出到 <数据目录>/outputs/发票报销汇总表.xlsx。
"""
import glob, os, re, sys, unicodedata
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.environ.get("INVOICE_WORKDIR") or os.getcwd()
OUT_DIR = os.path.join(BASE, "outputs")
os.makedirs(OUT_DIR, exist_ok=True)

def norm(t):
    return unicodedata.normalize("NFKC", t or "")

def extract(pdf):
    with pdfplumber.open(pdf) as d:
        text = norm("".join(pg.extract_text() or "" for pg in d.pages))
    if not text.strip():
        return None
    inv = {"file": os.path.basename(pdf), "via": "text"}
    m = re.search(r"发票号码[:：]\s*(\d{8,20})", text)
    inv["number"] = m.group(1) if m else None
    m = re.search(r"开票日期[:：]\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", text)
    inv["date"] = "%s-%02d-%02d" % (m.group(1), int(m.group(2)), int(m.group(3))) if m else None
    m = re.search(r"名称[:：](.+?)\s+名称[:：](.+)", text)
    inv["seller"] = m.group(2).strip() if m else None
    m = re.search(r"\*([^*\n]+)\*(\S+)", text)
    inv["item"] = ("*%s*%s" % (m.group(1), m.group(2))) if m else None
    m = re.search(r"合\s*计\s*[¥￥]([\d,.]+)\s*[¥￥]([\d,.]+)", text)
    inv["amount"] = float(m.group(1).replace(",", "")) if m else None
    inv["tax"] = float(m.group(2).replace(",", "")) if m else None
    m = re.search(r"[（(]小写[)）]\s*[¥￥]([\d,.]+)", text)
    inv["total"] = float(m.group(1).replace(",", "")) if m else None
    first = next((ln for ln in text.splitlines() if ln.strip()), "")
    inv["type"] = first.strip() if "发票" in first else "—"
    return inv

def extract_ocr(pdf):
    """百度云增值税发票识别兜底：用于扫描件/无文本层 PDF。无配置或失败返回 None。"""
    try:
        import baidu_ocr
        baidu_ocr.load_config()  # 未配置 Key 时直接放弃 OCR 路径
    except Exception:
        return None
    try:
        f = baidu_ocr.parse_fields(baidu_ocr.ocr_invoice_file(pdf))
    except Exception as e:
        print("OCR-ERR", os.path.basename(pdf), "->", e)
        return None
    if not f.get("invoice_number"):
        return None
    return {
        "file": os.path.basename(pdf),
        "via": "ocr",
        "number": f["invoice_number"],
        "date": f["invoice_date"] or None,
        "seller": f["seller_name"] or None,
        "item": f["item_description"] or None,
        "amount": f["amount"],
        "tax": f["tax"],
        "total": f["total"],
        "type": f["invoice_type"] or "—",
    }

pdfs = sorted(glob.glob(os.path.join(BASE, "invoices", "*.pdf")))
rows, review, dups = [], [], []
seen_numbers = set()  # 同一发票号码可能经附件与链接重复入库，或平台每次下载重新生成 PDF

def complete(inv):
    return bool(inv and inv.get("number") and inv.get("total") is not None and inv.get("seller"))

for p in pdfs:
    inv = None
    try:
        inv = extract(p)
    except Exception:
        inv = None
    if not complete(inv):
        ocr_inv = extract_ocr(p)  # 扫描件/文本层提取不全时走百度 OCR 兜底
        if complete(ocr_inv):
            inv = ocr_inv
    if not complete(inv):
        review.append((os.path.basename(p), "文本层与 OCR 均无法完整提取（可能为扫描件非发票或票面异常）"))
        continue
    ok = abs((inv["amount"] or 0) + (inv["tax"] or 0) - inv["total"]) <= 0.01
    if not ok:
        review.append((os.path.basename(p), "金额+税额≠价税合计"))
        continue
    if inv["number"] in seen_numbers:
        dups.append((os.path.basename(p), inv["number"]))
        continue
    seen_numbers.add(inv["number"])
    rows.append(inv)

rows.sort(key=lambda r: r["date"])

# ---------- 生成 xlsx ----------
wb = Workbook()
ws = wb.active
ws.title = "报销汇总"
headers = ["序号", "发票号码", "开票日期", "销售方名称", "项目摘要", "金额(元)", "税额(元)", "价税合计(元)", "发票类型"]
ncol = len(headers)
thin = Side(style="thin")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
t = ws.cell(row=1, column=1, value="发票报销汇总表")
t.font = Font(name="SimSun", size=14, bold=True)
t.alignment = Alignment(horizontal="center", vertical="center")

for c, h in enumerate(headers, 1):
    cell = ws.cell(row=2, column=c, value=h)
    cell.font = Font(name="SimSun", bold=True)
    cell.fill = PatternFill("solid", start_color="D9D9D9")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

for i, r in enumerate(rows):
    rr = 3 + i
    vals = [i + 1, r["number"], r["date"], r["seller"], r["item"], r["amount"], r["tax"], r["total"], r["type"]]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=rr, column=c, value=v)
        cell.font = Font(name="SimSun")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        if c in (6, 7, 8):
            cell.number_format = "0.00"

tr = 3 + len(rows)
ws.cell(row=tr, column=5, value="合计").font = Font(name="SimSun", bold=True)
dbl_top = Border(left=thin, right=thin, bottom=thin, top=Side(style="double"))
for c in range(1, ncol + 1):
    ws.cell(row=tr, column=c).border = dbl_top
    ws.cell(row=tr, column=c).font = Font(name="SimSun", bold=True)
    ws.cell(row=tr, column=c).alignment = Alignment(horizontal="center")
for c, col in ((6, "F"), (7, "G"), (8, "H")):
    ws.cell(row=tr, column=c, value="=SUM(%s3:%s%d)" % (col, col, tr - 1)).number_format = "0.00"

widths = [6, 24, 12, 34, 30, 12, 12, 14, 24]
for c, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(c)].width = max(w, 12)

out = os.path.join(OUT_DIR, "发票报销汇总表.xlsx")
wb.save(out)

print("EXTRACTED %d invoices:" % len(rows))
for r in rows:
    print(" ", "[%s]" % r.get("via", "text"), r["date"], r["number"], r["seller"], r["amount"], r["tax"], r["total"])
print("REVIEW:")
for name, why in review:
    print(" ", name, "->", why)
print("PY-CHECK totals:", round(sum(r["amount"] for r in rows), 2), round(sum(r["tax"] for r in rows), 2), round(sum(r["total"] for r in rows), 2))
print("SAVED", out)
