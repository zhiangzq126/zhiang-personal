import logging
from pathlib import Path
from typing import Iterable, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

SHEET_NAME = "发票汇总"
MONEY_FORMAT = "#,##0.00"

COLUMNS = [
    ("邮件日期", 16, "mail_date"),
    ("邮件主题", 34, "mail_subject"),
    ("发件人", 22, "mail_sender"),
    ("发票类型", 18, "invoice_type"),
    ("发票代码", 14, "code"),
    ("发票号码", 22, "number"),
    ("开票日期", 12, "date"),
    ("购买方名称", 26, "buyer_name"),
    ("购买方税号", 20, "buyer_tax_id"),
    ("销售方名称", 26, "seller_name"),
    ("销售方税号", 20, "seller_tax_id"),
    ("不含税金额", 12, "amount"),
    ("税额", 10, "tax"),
    ("价税合计", 12, "total"),
    ("来源", 10, "source"),
    ("文件", 30, "_file_link"),
    ("状态", 14, "status"),
    ("备注", 30, "note"),
]
MONEY_KEYS = {"amount", "tax", "total"}
KEY_COLUMN_INDEX = next(i for i, col in enumerate(COLUMNS, start=1) if col[2] == "number")
FILE_COLUMN_INDEX = next(i for i, col in enumerate(COLUMNS, start=1) if col[2] == "_file_link")


def _ensure_sheet(path: Path):
    if path.exists():
        book = load_workbook(path)
        sheet = book[SHEET_NAME] if SHEET_NAME in book.sheetnames else book.create_sheet(SHEET_NAME)
        if sheet.max_row < 1 or sheet["A1"].value is None:
            _write_header(sheet)
        return book, sheet

    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_NAME
    _write_header(sheet)
    return book, sheet


def _write_header(sheet) -> None:
    fill = PatternFill("solid", fgColor="DDEBF7")
    for index, (title, width, _) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def existing_keys(path: Path) -> set:
    """已入表的发票，用发票号或文件 SHA1 去重。"""
    if not Path(path).exists():
        return set()
    book = load_workbook(path, read_only=True)
    if SHEET_NAME not in book.sheetnames:
        return set()
    keys = set()
    for row in book[SHEET_NAME].iter_rows(min_row=2, values_only=True):
        number = row[KEY_COLUMN_INDEX - 1]
        if number:
            keys.add(f"NO:{number}")
        note = row[len(COLUMNS) - 1]
        if isinstance(note, str) and "SHA1:" in note:
            keys.add(note[note.index("SHA1:") : note.index("SHA1:") + 45])
    book.close()
    return keys


def append_invoices(path: Path, invoices: Iterable) -> int:
    invoices = list(invoices)
    if not invoices:
        return 0

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book, sheet = _ensure_sheet(path)
    row = sheet.max_row + 1

    for invoice in invoices:
        for index, (_, _, key) in enumerate(COLUMNS, start=1):
            if key == "_file_link":
                cell = sheet.cell(row=row, column=index, value=invoice.file_path.name)
                cell.hyperlink = invoice.file_path.resolve().as_uri()
                cell.font = Font(color="0563C1", underline="single")
                continue
            value = getattr(invoice, key, "")
            if key == "note" and not invoice.number:
                value = f"{value}；SHA1:{invoice.file_sha1}".strip("；")
            cell = sheet.cell(row=row, column=index, value=value)
            if key in MONEY_KEYS:
                cell.number_format = MONEY_FORMAT
        row += 1

    book.save(path)
    log.info("写入 %d 条到 %s", len(invoices), path)
    return len(invoices)
